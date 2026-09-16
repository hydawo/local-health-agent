"""Turn a container into the text the notes pipeline already ingests.

One job: `.docx` and image files come in, plain text goes out. Nothing here
touches the database, frontmatter, or chunking; `notes.py` does all of that
exactly as it does for a markdown file. Keeping the readers separate means a
new container is a new function here, not a new branch through `ingest_note`.

**Photos are evidence you can search, never data you can trend.** An image
is OCR'd and stored as searchable text; it is never parsed for lab values.
Tesseract on a flatbed scan is already the tool's weakest extraction path
(units and reference ranges measurably degrade), and a phone photo, with
skew, uneven light, and the camera's own sharpening, is worse again. A wrong
lab value silently poisons a trend line; a report someone wants in a trend
has to arrive as a PDF.
"""

from __future__ import annotations

from pathlib import Path

from ..logging_setup import get_logger
from . import records

log = get_logger("ingest.readers")

IMAGE_SUFFIXES: frozenset[str] = frozenset({".png", ".jpg", ".jpeg", ".heic"})
DOCX_SUFFIXES: frozenset[str] = frozenset({".docx"})

# python-docx exposes a heading's level only through its style name.
_HEADING_STYLE = "Heading "
_MAX_HEADING_LEVEL = 6


class HeicUnsupported(RuntimeError):
    """`.heic` needs `pillow-heif`, which lives in the `ocr` extra."""


def _heic_supported() -> bool:
    try:
        import pillow_heif
    except ImportError:
        return False
    pillow_heif.register_heif_opener()
    return True


def read_docx(path: Path) -> str:
    """Body text as markdown-ish plain text, in document order.

    Headings become `#` lines at their level so the notes section splitter
    gives chunks a heading trail; tables are flattened one row per line with
    ` | ` between cells, because a medication list is the likeliest thing to
    be in one and it must land under its heading, not at the end. Headers,
    footers, comments, tracked changes, and embedded images are ignored.
    """
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    lines: list[str] = []
    # Walk the body's own children rather than `document.paragraphs` then
    # `document.tables`: that pair loses the interleaving, which is the whole
    # point of a table sitting under a heading.
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            if style.startswith(_HEADING_STYLE):
                try:
                    level = int(style[len(_HEADING_STYLE):])
                except ValueError:
                    level = 1
                level = max(1, min(level, _MAX_HEADING_LEVEL))
                lines.append(f"{'#' * level} {text}")
            else:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    return "\n\n".join(lines)


def read_image(path: Path) -> str:
    """OCR a photo or screenshot. Plain text only.

    EXIF orientation is applied first: phones record rotation as metadata
    and Tesseract reads the raw pixels, so a sideways photo would otherwise
    OCR to nothing usable. Raises `records.OcrUnavailable` when Tesseract or
    pytesseract is missing and `HeicUnsupported` for `.heic` without
    `pillow-heif`; the caller reports both rather than skipping silently.
    """
    if path.suffix.lower() == ".heic" and not _heic_supported():
        raise HeicUnsupported(
            "reading .heic needs pillow-heif: pip install -e '.[ocr]'")

    from PIL import Image, ImageOps

    with Image.open(path) as opened:
        image = ImageOps.exif_transpose(opened)
        image.load()
    plain, _layout = records.ocr_image(image)
    return plain


__all__ = ["DOCX_SUFFIXES", "IMAGE_SUFFIXES", "HeicUnsupported",
           "read_docx", "read_image"]
